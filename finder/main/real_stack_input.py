"""Read workbook-labelled freeze threads; preserve evidence without inferring root causes."""

import argparse
from collections import Counter
import json
from pathlib import Path
import re
import sys

import openpyxl


THREAD = re.compile(r"^\[(?P<name>[^\]]+)\]\s+(?P<state>[A-Z_]+)(?:\s+(?P<details>.*))?$")
FRAME = re.compile(r"^(?:(topStack|problemModuleStack):\s*)?(\d+)\s+(.+)$")
LOCATION = re.compile(r"^(.*)\(([^()]*)\)$")
SOURCE = re.compile(r"^(.+\.(?:java|kt)):(\d+)$")
SEPARATOR = "cause freeze thread:"


def parse_stack(text):
  """Return all observed threads and labelled problem frames, including unparsed lines."""
  threads, diagnostics, unparsed = [], [], []
  current = None
  role = "blocked"
  separator_count = 0
  for original in text.splitlines():
    line = original.strip()
    if not line or re.fullmatch(r"=+", line):
      continue
    if line.startswith(SEPARATOR):
      separator_count += 1
      role = "cause"
      current = None
      line = line[len(SEPARATOR):].strip()
      if not line:
        continue
    header = THREAD.fullmatch(line)
    if header:
      current = {**header.groupdict(), "role": role,
                 "is_ui_thread": bool(re.fullmatch(r"AWT-EventQueue-\d+", header["name"])),
                 "frames": [], "unparsed_lines": []}
      threads.append(current)
      continue
    match = FRAME.fullmatch(line)
    if current is not None and match:
      label, index, body = match.groups()
      location = LOCATION.fullmatch(body)
      symbol, position = location.groups() if location else (None, None)
      source = SOURCE.fullmatch(position or "")
      current["frames"].append({"index": int(index), "label": label, "symbol": symbol,
                                "location": position, "file": source[1] if source else None,
                                "line": int(source[2]) if source else None,
                                "elided": body == "*", "raw": original})
      if location is None and body != "*":
        diagnostics.append("UNRECOGNIZED_FRAME")
    else:
      (current["unparsed_lines"] if current else unparsed).append(original)
      diagnostics.append("UNRECOGNIZED_LINE")
  causes = [thread for thread in threads if thread["role"] == "cause"]
  ui_threads = [thread for thread in threads if thread["is_ui_thread"]]
  top = [frame for thread in causes for frame in thread["frames"] if frame["label"] == "topStack"]
  problem = [frame for thread in causes for frame in thread["frames"] if frame["label"] == "problemModuleStack"]
  if separator_count != 1:
    diagnostics.append("CAUSE_SEPARATOR_COUNT")
  if len(causes) != 1:
    diagnostics.append("CAUSE_THREAD_COUNT")
  if not ui_threads:
    diagnostics.append("UI_THREAD_NOT_PRESENT")
  if any(not thread["frames"] for thread in threads):
    diagnostics.append("EMPTY_THREAD_STACK")
  if len(top) != 1:
    diagnostics.append("TOP_STACK_COUNT")
  if len(problem) != 1:
    diagnostics.append("PROBLEM_MODULE_STACK_COUNT")
  return {"status": "PARSED" if not diagnostics else "PARTIAL" if threads else "FAILED",
          "threads": threads, "ui_threads": ui_threads, "cause_threads": causes,
          "problem": {"evidence_kind": "WORKBOOK_LABELS", "top_stack": top,
                      "problem_module_stack": problem, "root_cause": None},
          "diagnostics": list(dict.fromkeys(diagnostics)), "unparsed_lines": unparsed, "raw_stack": text}


def read_records(workbook, sheet_name="问题线程堆栈"):
  """Find columns by header, preserving each nonempty input row and Excel row number."""
  book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
  try:
    if sheet_name not in book.sheetnames:
      raise ValueError(f"Missing worksheet: {sheet_name}")
    rows = enumerate(book[sheet_name].iter_rows(values_only=True), 1)
    for header_row, values in rows:
      headers = [value.strip() if isinstance(value, str) else value for value in values]
      if "问题线程堆栈" in headers:
        if headers.count("问题线程堆栈") != 1:
          raise ValueError("Duplicate stack column")
        stack_column = headers.index("问题线程堆栈")
        count_column = headers.index("堆栈出现次数") if "堆栈出现次数" in headers else None
        break
    else:
      raise ValueError("Missing column: 问题线程堆栈")
    for row, values in rows:
      value = values[stack_column] if stack_column < len(values) else None
      if value is None or value == "":
        continue
      # This workbook splits long stacks at Excel's 32,767-character cell limit.
      column, part = stack_column + 1, value
      while isinstance(part, str) and len(part) == 32767 and column < len(values):
        if headers[column] is not None and headers[column] != "":
          break
        part = values[column]
        if not isinstance(part, str):
          break
        value += part
        column += 1
      occurrences = values[count_column] if count_column is not None and count_column < len(values) else None
      record = parse_stack(value if isinstance(value, str) else str(value))
      if not isinstance(value, str):
        record["status"] = "FAILED"
        record["diagnostics"].append("STACK_NOT_TEXT")
      yield {"record_id": f"workbook-row-{row}", "sheet": sheet_name, "row": row,
             "occurrences": occurrences, **record}
  finally:
    book.close()


def main(argv=None):
  parser = argparse.ArgumentParser(description=__doc__)
  parser.add_argument("--workbook", required=True, type=Path)
  parser.add_argument("--sheet", default="问题线程堆栈")
  parser.add_argument("--output", type=Path, help="JSONL destination; defaults to stdout")
  args = parser.parse_args(argv)
  if args.output and args.output.resolve() == args.workbook.resolve():
    parser.error("Output must not overwrite the workbook")
  counts = Counter()
  output = None
  try:
    # Exclusive creation protects both existing results and arbitrary user files.
    output = args.output.open("x", encoding="utf-8") if args.output else sys.stdout
    for record in read_records(args.workbook, args.sheet):
      output.write(json.dumps(record, ensure_ascii=False) + "\n")
      counts[record["status"]] += 1
  finally:
    if output is not None and output is not sys.stdout:
      output.close()
  print(json.dumps({"records": sum(counts.values()), "statuses": dict(counts)}), file=sys.stderr)
  return 0


if __name__ == "__main__":
  raise SystemExit(main())
